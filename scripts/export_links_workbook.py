"""Classeur des liens de référence des concours (Y-Seck-liens-references-concours.xlsx), mis à jour depuis le catalogue.

Les trois premières feuilles (référentiel des concours, organisateurs, calendrier) sont reconstruites à partir de
data/catalogue ; les quatre suivantes (annales, bibliographie, prépas, fiches métiers) sont reprises du classeur
précédent, leurs données n'ayant pas changé.
Usage : python3 scripts/export_links_workbook.py [sortie.xlsx] [classeur précédent.xlsx]
"""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / 'data/research/2026-10-22/Y-Seck-liens-references-concours.xlsx'
FPS = {'E': 'État', 'T': 'Territoriale', 'H': 'Hospitalière'}


def catalogue():
    manifest = json.loads((ROOT / 'data/catalogue.json').read_text(encoding='utf-8'))
    db = {}
    for part in manifest['parts']:
        for key, value in json.loads((ROOT / 'data' / part).read_text(encoding='utf-8')).items():
            if isinstance(value, list):
                db.setdefault(key, []).extend(value)
            else:
                db[key] = value
    return db


def versants(fp):
    return ', '.join(FPS[f] for f in fp if f in FPS)


def day(iso):
    return '-'.join(reversed(iso.split('-'))) if iso else None


def referentiel(db):
    rows = []
    for c in db['concours']:
        head = [c['name'], c['cat'], versants(c['fp'])]
        rows.append(head + ['Page officielle', c['source']])
        if c.get('annales'):
            years = c['annales'].get('years')
            rows.append(head + ['Annales' + (' (' + years + ')' if years else ''), c['annales']['url']])
        for s in c.get('sessions', []):
            if s.get('url') and s['url'] != c['source']:
                rows.append(head + ['Session ' + s['session'], s['url']])
        for n in c.get('next', []):
            if n.get('source'):
                rows.append(head + ['Évoluer : ' + n['toName'], n['source']])
    return rows


def organisateurs(db):
    rows = []
    for r in db['routes']:
        if r.get('source'):
            rows.append([r['name'], r['cat'], versants(r['fp']), None, 'Page du corps', r['source']])
    cats = {r['id']: r['cat'] for r in db['routes']}
    for e in db.get('examDetails', []):
        head = [e['name'], cats.get(e['routeId'], ''), versants(e['fp']), e.get('org')]
        if e.get('source'):
            rows.append(head + ['Source', e['source']])
        for link in e.get('officialLinks', []):
            rows.append(head + [link['label'], link['url']])
    return sorted(rows, key=lambda r: (r[0].lower(), r[4] != 'Page du corps'))


def calendrier(db):
    rows = []
    for c in db['calendar']:
        if c.get('duplicateOf'):
            continue
        rows.append([c['name'], c['cat'], versants(c['fp']), c.get('voie'), c.get('openLabel') or day(c.get('open')),
                     c.get('closeLabel') or day(c.get('date')), c.get('source'), c.get('url')])
    return sorted(rows, key=lambda r: r[0].lower())


SHEETS = [
    ('Référentiel concours', ['Concours', 'Catégorie', 'Versant', 'Type de lien', 'Lien'], [48, 10, 14, 34, 90], referentiel),
    ('Concours et organisateurs', ['Concours ou corps', 'Catégorie', 'Versant', 'Organisé par', 'Type de lien', 'Lien'], [50, 10, 14, 40, 30, 90], organisateurs),
    ('Calendrier des sessions', ['Concours', 'Catégorie', 'Versant', 'Voie', 'Ouverture', 'Clôture', 'Source', 'Lien'], [46, 10, 14, 24, 30, 30, 14, 90], calendrier),
]


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'work/Y-Seck-liens-references-concours.xlsx'
    template = Path(sys.argv[2]) if len(sys.argv) > 2 else TEMPLATE
    db = catalogue()
    wb = openpyxl.load_workbook(template)
    header = {'font': Font(bold=True, color='FFFFFF'), 'fill': PatternFill('solid', fgColor='2B04CC'), 'alignment': Alignment(vertical='center', wrap_text=True)}
    for index, (title, headers, widths, build) in enumerate(SHEETS):
        del wb[title]
        ws = wb.create_sheet(title, index)
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill, cell.alignment = header['font'], header['fill'], header['alignment']
        for row in build(db):
            ws.append(row)
            link = ws.cell(ws.max_row, len(headers))
            if isinstance(link.value, str) and link.value.startswith('http'):
                link.hyperlink = link.value
                link.style = 'Hyperlink'
        for i, width in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = width
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(out, {title: wb[title].max_row - 1 for title, *_ in SHEETS})


if __name__ == '__main__':
    main()
